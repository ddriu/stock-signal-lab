from dataclasses import asdict
from io import BytesIO
import json

from openpyxl import load_workbook
import numpy as np
import pandas as pd
import pytest

from config import BacktestConfig, StrategyConfig
from src.analysis_policy import encode_policy_note
from src.backtesting import run_backtest
from src.indicators import add_indicators
from src.model_results_engine import ExecutionCosts
from src.model_results_service import _legacy_reference, build_model_results, export_model_results_excel, public_frame, report_tables
from src.signal_engine import add_signal_columns


def inputs():
    days = pd.bdate_range("2025-01-02", periods=8)
    frame = pd.DataFrame(dict(open=100.0, high=101.0, low=99.0, close=100.0,
        fx_open=1.0, fx_close=1.0, reduce_setup=False, sell_setup=False, sma_long=1.0), index=days)
    frame["open_at"] = pd.to_datetime(days, utc=True) + pd.Timedelta(hours=14, minutes=30)
    frame["close_at"] = pd.to_datetime(days, utc=True) + pd.Timedelta(hours=21)
    frame.attrs.update(display_currency="USD", strategy_config=asdict(StrategyConfig()),
        session_schedule=[dict(session=day.date().isoformat(), open_at=row.open_at.isoformat(),
            close_at=row.close_at.isoformat()) for day, row in frame.iterrows()],
        expected_sessions=[day.date().isoformat() for day in days],
        required_last_close_at=frame.iloc[-1].close_at.isoformat())
    snapshots = pd.DataFrame([dict(id=index + 1, ticker="SYNTHETIC", analyzed_at=days[index].date().isoformat(),
        created_at=(frame.iloc[index].close_at + pd.Timedelta(minutes=1)).isoformat(),
        entry_label="Esperar" if index == 0 else "Entrada fuerte", entry_score=60 if index == 0 else 86,
        sector="=1+1", owner="private@example.invalid", note="") for index in range(3)])
    bars = {ticker: frame.copy(deep=True) for ticker in ("SYNTHETIC", "SPY", "QQQ", "VT")}
    return snapshots, bars


def build(**kwargs):
    snapshots, bars = inputs()
    return build_model_results(snapshots, bars, start="2025-01-02", end="2025-01-13", **kwargs)


def test_builds_eight_results_and_twenty_four_separate_flow_matched_references():
    report = build()
    assert list(report.theoretical) == list(report.budget) == ["A", "B", "C", "D"]
    assert len(report.benchmarks) == 24
    assert set(report.benchmarks.benchmark) == {"SPY", "QQQ", "VT"}
    assert report.manifest["prospective_status"].startswith("No activado")
    assert report.original_backtest.status.eq("N/D").all()
    assert report.original_reference.manifest["config"]["exit_policy"] == "original"
    assert report.original_reference.manifest["config"]["amount_usd"] == 100


def test_budget_benchmarks_receive_one_initial_flow_not_recycled_buys():
    snapshots, bars = inputs()
    frame = bars["SYNTHETIC"]
    frame.loc[frame.index[2], "reduce_setup"] = True
    snapshots.loc[2, "analyzed_at"] = frame.index[3].date().isoformat()
    snapshots.loc[2, "created_at"] = (frame.iloc[3].close_at + pd.Timedelta(minutes=1)).isoformat()
    report = build_model_results(snapshots, bars, start="2025-01-02", end="2025-01-13",
                                budget_usd=100, costs=ExecutionCosts(0, 0, 0))
    assert report.budget["A"].metrics["invested_notional_usd"] == 200
    budget_refs = report.benchmarks.loc[report.benchmarks["mode"].eq("budget")]
    assert budget_refs.external_contributions_usd.eq(100).all()
    assert budget_refs.value_usd.eq(100).all()
    theoretical_a = report.benchmarks.loc[report.benchmarks["mode"].eq("theoretical") & report.benchmarks.model.eq("A")]
    assert theoretical_a.external_contributions_usd.eq(200).all()


def test_budget_benchmark_costs_never_overspend_initial_cash():
    report = build()
    refs = report.benchmarks.loc[report.benchmarks["mode"].eq("budget")]
    assert refs.data_complete.all()
    expected_notional = (6000 - 1) / 1.001
    assert refs.invested_notional_usd.tolist() == pytest.approx([expected_notional] * 12)
    assert refs.value_usd.tolist() == pytest.approx([expected_notional] * 12)
    assert refs.fees_usd.tolist() == pytest.approx([6000 - expected_notional] * 12)
    assert refs.roi_pct.lt(0).all()


def test_missing_benchmark_is_unknown_not_zero():
    snapshots, bars = inputs()
    del bars["QQQ"]
    report = build_model_results(snapshots, bars, start="2025-01-02", end="2025-01-13")
    refs = report.benchmarks.loc[report.benchmarks.benchmark.eq("QQQ")]
    assert refs.value_usd.isna().all()
    assert not refs.data_complete.any()
    assert refs.excess_return_pp.isna().all()
    assert refs.outperformed_benchmark.isna().all()


def test_growth_source_is_not_a_technical_fallback_or_score_mix():
    snapshots, bars = inputs()
    empty = build_model_results(snapshots, bars, start="2025-01-02", end="2025-01-13", source="growth")
    assert empty.signals.empty
    assert all(result.lots.empty for result in empty.theoretical.values())
    for index in snapshots.index:
        snapshots.loc[index, "note"] = encode_policy_note("", dict(policy_version="synthetic_v1",
            ticker="SYNTHETIC", primary=dict(strategy_id="growth_momentum", label="Growth",
            status="Esperar mejor precio" if index == 0 else "Entrada fuerte", score=84,
            data_status="complete", entry_candidate=index != 0), comparative={}, position={}))
    growth = build_model_results(snapshots, bars, start="2025-01-02", end="2025-01-13", source="growth")
    assert not growth.theoretical["B"].lots.empty
    assert growth.theoretical["C"].lots.empty
    assert set(growth.signals.entry_score) == {84}
    assert set(growth.signals.technical_entry_score) == {60, 86}


def _growth_cohort_inputs():
    _, bars = inputs()
    frame = bars["SYNTHETIC"]
    snapshots = []
    for index in range(5):
        day = frame.index[index].date().isoformat()
        primary = dict(strategy_id="growth_momentum", label="Crecimiento + momentum",
            status="Vigilancia activa" if index == 0 else "Entrada fuerte", score=88,
            data_status="complete", entry_candidate=index != 0, as_of=day)
        policy = dict(policy_version="dual_strategy_v1", ticker="SYNTHETIC", as_of=day,
            primary=primary, comparative={}, position={})
        if index != 3:
            policy["data_quality"] = dict(version="daily_data_quality_v1", as_of=day,
                status="review", block_entry=index == 1, metric_periods={})
        snapshots.append(dict(id=index + 1, ticker="SYNTHETIC", analyzed_at=day,
            created_at=(frame.iloc[index].close_at + pd.Timedelta(minutes=1)).isoformat(),
            entry_label="Entrada fuerte", entry_score=90, price=100.0,
            note=encode_policy_note("", policy)))
    return pd.DataFrame(snapshots), bars


def test_recorded_verified_growth_cohort_has_own_entries_without_changing_raw_original_reference():
    snapshots, bars = _growth_cohort_inputs()
    kwargs = dict(start="2025-01-02", end="2025-01-13", source="growth",
                  costs=ExecutionCosts(0, 0, 0))
    raw = build_model_results(snapshots, bars, **kwargs)
    raw_explicit = build_model_results(snapshots, bars, cohort="raw", **kwargs)
    verified = build_model_results(snapshots, bars, cohort="verified", **kwargs)
    assert raw.manifest == raw_explicit.manifest
    assert raw.theoretical["A"].metrics["executed_entries"] == 4
    assert verified.theoretical["A"].metrics["executed_entries"] == 2
    # Ineligible observations must remain in the activation state machine.
    assert raw.theoretical["B"].metrics["executed_entries"] == 1
    assert verified.theoretical["B"].metrics["executed_entries"] == 2
    assert verified.signals["entry_label"].tolist() == ["Vigilancia activa",
        "Sin entrada verificada", "Entrada fuerte", "Sin entrada verificada", "Entrada fuerte"]
    assert verified.observations.coverage["blocked_entry_observations"] == 2
    assert verified.manifest["cohort"] == "verified"
    assert "cohort" not in raw.manifest
    assert raw.manifest["report_hash"] != verified.manifest["report_hash"]
    assert verified.manifest["original_reference_cohort"] == "raw"
    pd.testing.assert_frame_equal(raw.original_reference.lots, verified.original_reference.lots)
    for model in verified.theoretical.values():
        assert all(row.entry_at > row.observed_at for row in model.lots.itertuples())
    assert not set(raw.theoretical["A"].lots.event_id) & set(verified.theoretical["A"].lots.event_id)


def test_late_same_session_verification_cannot_retroactively_create_an_executable_entry():
    snapshots, bars = _growth_cohort_inputs()
    # Keep only a blocked first strong observation; the later same-session
    # revision is even available before the next open, but this conservative
    # cohort deliberately uses the identical first-observation rule as raw.
    first = snapshots.iloc[:2].copy()
    later = first.iloc[-1].to_dict()
    later["id"] = 99
    later["created_at"] = (pd.Timestamp(later["created_at"]) + pd.Timedelta(hours=1)).isoformat()
    from src.analysis_policy import decode_policy_note
    policy = decode_policy_note(later["note"])
    policy["data_quality"]["block_entry"] = False
    later["note"] = encode_policy_note("", policy)
    kwargs = dict(start="2025-01-02", end="2025-01-13", source="growth", cohort="verified")
    before = build_model_results(first, bars, **kwargs)
    after = build_model_results(pd.concat([first, pd.DataFrame([later])], ignore_index=True), bars, **kwargs)
    pd.testing.assert_frame_equal(before.signals, after.signals)
    assert all(model.lots.empty for model in after.theoretical.values())
    assert after.observations.rejected.iloc[0]["reason"] == "later_revision"
    assert after.benchmarks.outperformed_benchmark.isna().all()


def test_unknown_quality_cohort_cannot_be_used_with_the_original_technical_source():
    with pytest.raises(ValueError, match="cohort"):
        build(source="technical", cohort="verified")


def test_mismatched_config_or_obsolete_preparation_rejected():
    snapshots, bars = inputs()
    bars["SYNTHETIC"].attrs["strategy_config"]["sma_long"] = 201
    with pytest.raises(ValueError, match="configuration mismatch"):
        build_model_results(snapshots, bars, start="2025-01-02", end="2025-01-13")
    snapshots, bars = inputs()
    bars["SYNTHETIC"].attrs["audit_data_version"] = "strategy-audit-data-v2"
    with pytest.raises(ValueError, match="version mismatch"):
        build_model_results(snapshots, bars, start="2025-01-02", end="2025-01-13")


def test_full_excel_is_in_memory_formula_safe_and_identifier_redacted():
    report = build()
    content = export_model_results_excel(report)
    assert content[:2] == b"PK"
    workbook = load_workbook(BytesIO(content), data_only=False)
    assert {"Resumen", "Evolución", "Lotes", "Ventas", "Atribución", "Índices", "Reglas y procedencia"}.issubset(workbook.sheetnames)
    formula = next(cell for row in workbook["Señales"] for cell in row if cell.value == "=1+1")
    assert formula.data_type == "s"
    text = " ".join(str(cell.value) for sheet in workbook for row in sheet for cell in row)
    assert "private@example.invalid" not in text
    encoded = "".join(str(row[1]) for row in workbook["Reglas y procedencia"].iter_rows(min_row=2, values_only=True))
    assert json.loads(encoded)["report_hash"] == report.manifest["report_hash"]
    assert len(report_tables(report)["Resumen"]) == 8


def test_public_frame_does_not_mutate_or_expose_nested_account_identity():
    source = pd.DataFrame([dict(owner="private@example.invalid", note="private@example.invalid",
        raw_record=dict(email="private@example.invalid", safe="retained"))])
    visible = public_frame(source)
    assert "owner" not in visible
    assert "private@example.invalid" not in repr(visible.to_dict("records"))
    assert "private@example.invalid" in source.note.iloc[0]


def test_legacy_reference_rebuilds_original_indicators_before_cutting_230_session_window():
    days = pd.bdate_range("2024-01-02", periods=490)
    step = np.arange(len(days))
    close = 100 + step * .08 + 2.5 * np.sin(step / 5) + np.sin(step / 17)
    opened = close + .2 * np.sin(step / 3)
    frame = pd.DataFrame(dict(open=opened, close=close,
        high=np.maximum(opened, close) + .7, low=np.minimum(opened, close) - .7,
        volume=1_000_000 + 100_000 * np.sin(step / 7)), index=days)
    # Shape delivered by prepare_audit_bars: no short/medium SMA, RSI, MACD
    # or scoring columns; an existing sma_long must not be blindly trusted.
    prepared = frame.assign(sma_long=-999, reduce_setup=False, sell_setup=False)
    prepared.attrs["display_currency"] = "GBP"
    prepared["close_at"] = pd.to_datetime(days, utc=True) + pd.Timedelta(hours=21)
    start, end = days[250].date().isoformat(), days[479].date().isoformat()
    strategy, settings = StrategyConfig(), BacktestConfig()
    causal = frame.loc[:days[479]].copy()
    direct_input = add_signal_columns(add_indicators(causal, strategy), strategy).loc[days[250]:]
    assert len(direct_input) == 230
    direct = run_backtest(direct_input, strategy, settings)
    assert direct.metrics["completed_trades"] > 0
    actual = _legacy_reference({"SYNTHETIC": prepared}, start, end).iloc[0]
    assert actual.status == "Calculado"
    assert actual.currency == "GBP"
    assert actual.capital_units == "moneda nativa; sin conversión FX"
    assert actual.initial_capital_native == settings.initial_capital
    for key, value in direct.metrics.items():
        assert actual[key] == pytest.approx(value)
    assert prepared.sma_long.eq(-999).all()
    # Extreme prices after the cutoff cannot alter the reference or warmup.
    future = prepared.copy(deep=True)
    future.loc[days[480]:, ["open", "high", "low", "close"]] = [10000, 11000, 9000, 10000]
    later = _legacy_reference({"SYNTHETIC": future}, start, end)
    pd.testing.assert_frame_equal(_legacy_reference({"SYNTHETIC": prepared}, start, end), later)


def test_legacy_reference_keeps_original_minimum_202_sessions_despite_warmup():
    days = pd.bdate_range("2024-01-02", periods=300)
    frame = pd.DataFrame(dict(open=100.0, high=101.0, low=99.0, close=100.0,
                             volume=1_000_000), index=days)
    reference = _legacy_reference({"SYNTHETIC": frame}, days[-10].date().isoformat(),
                                  days[-1].date().isoformat()).iloc[0]
    assert reference.status == "N/D"
    assert "202 sesiones" in reference.reason


def test_benchmark_outperformance_is_explicit_only_for_complete_nonempty_cohorts():
    report = build()
    refs = report.benchmarks.loc[report.benchmarks["mode"].eq("budget") & report.benchmarks.model.eq("A")]
    model_return = report.budget["A"].metrics["net_return_pct"]
    assert report.budget["A"].metrics["coverage_complete"]
    assert refs.model_coverage_complete.all()
    assert refs.model_net_return_pct.tolist() == pytest.approx([model_return] * 3)
    assert refs.excess_return_pp.tolist() == pytest.approx((model_return - refs.roi_pct).tolist())
    assert refs.outperformed_benchmark.eq(True).all()
    snapshots, bars = inputs()
    snapshots.loc[1, "created_at"] = (bars["SYNTHETIC"].iloc[1].close_at - pd.Timedelta(hours=1)).isoformat()
    incomplete = build_model_results(snapshots, bars, start="2025-01-02", end="2025-01-13")
    partial = incomplete.benchmarks.loc[incomplete.benchmarks.model.eq("A")]
    assert not partial.model_coverage_complete.any()
    assert partial.model_net_return_pct.notna().all()  # readable, explicitly partial
    assert partial.excess_return_pp.isna().all()
    assert partial.outperformed_benchmark.isna().all()
    empty = build_model_results(snapshots, bars, start="2025-01-02", end="2025-01-13", source="growth")
    assert empty.benchmarks.model_net_return_pct.isna().all()
    assert empty.benchmarks.excess_return_pp.isna().all()
    assert empty.benchmarks.outperformed_benchmark.isna().all()


def test_each_model_has_its_own_exact_matched_hold_not_the_a_cohort():
    snapshots, bars = inputs()
    snapshots.loc[1, "entry_score"] = 84
    snapshots.loc[2, ["entry_label", "entry_score"]] = ["Esperar", 60]
    frame = bars["SYNTHETIC"]
    additional = []
    for index, score in ((3, 86), (4, 99)):
        row = snapshots.iloc[1].to_dict()
        row.update(id=index + 1, analyzed_at=frame.index[index].date().isoformat(),
            created_at=(frame.iloc[index].close_at + pd.Timedelta(minutes=1)).isoformat(), entry_score=score)
        additional.append(row)
    snapshots = pd.concat([snapshots, pd.DataFrame(additional)], ignore_index=True)
    report = build_model_results(snapshots, bars, start="2025-01-02", end="2025-01-13")
    assert [len(report.theoretical[key].lots) for key in ("A", "B", "C")] == [3, 2, 1]
    assert len(report.matched_hold_summary) == 6
    for mode, models in (("theoretical", report.theoretical), ("budget", report.budget)):
        for key in ("A", "B", "C"):
            held = report.matched_hold[mode][key]
            assert held.metrics["cohort_matches_model"]
            assert held.metrics["budget_feasible"]
            columns = ["event_id", "entry_at", "quantity", "entry_price_native", "entry_fx", "entry_debit_usd"]
            expected = models[key].lots.loc[:, columns].sort_values("event_id").reset_index(drop=True)
            actual = held.lots.loc[:, columns].sort_values("event_id").reset_index(drop=True)
            pd.testing.assert_frame_equal(expected, actual)
            assert held.trades.empty
    exported = report_tables(report)
    assert {"Mantener mismos lotes", "Lotes mantener", "Evolución mantener"}.issubset(exported)
    workbook = load_workbook(BytesIO(export_model_results_excel(report)))
    assert "Mantener mismos lotes" in workbook.sheetnames


def test_exact_hold_budget_shortfall_is_nd_not_credit_or_partial_cohort():
    snapshots, bars = inputs()
    frame = bars["SYNTHETIC"]
    frame.loc[frame.index[2], "reduce_setup"] = True
    snapshots.loc[2, "analyzed_at"] = frame.index[3].date().isoformat()
    snapshots.loc[2, "created_at"] = (frame.iloc[3].close_at + pd.Timedelta(minutes=1)).isoformat()
    report = build_model_results(snapshots, bars, start="2025-01-02", end="2025-01-13",
                                budget_usd=100, costs=ExecutionCosts(0, 0, 0))
    assert len(report.budget["A"].lots) == 2
    held = report.matched_hold["budget"]["A"]
    assert held.metrics["cohort_matches_model"]
    assert not held.metrics["budget_feasible"]
    assert held.metrics["capital_required_usd"] == 200
    assert held.metrics["cash_shortfall_usd"] == 100
    assert held.metrics["final_value_usd"] is None
    assert held.metrics["net_return_pct"] is None
    assert held.metrics["cash_usd"] is None
    assert held.metrics["reference_executed_entries"] == 0
    assert held.daily.cash_usd.isna().all()
    assert held.lots.reference_execution_status.eq("not_funded").all()
    assert report.matched_hold["budget"]["B"].metrics["budget_feasible"]
    assert report.matched_hold["budget"]["C"].metrics["budget_feasible"]
    row = report.matched_hold_summary.loc[report.matched_hold_summary["mode"].eq("budget") &
        report.matched_hold_summary.model.eq("A")].iloc[0]
    assert pd.isna(row.excess_return_pp)
    assert pd.isna(row.hold_net_return_pct)
