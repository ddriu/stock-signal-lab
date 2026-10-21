from dataclasses import replace
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.strategy_audit_engine import AuditConfig, run_strategy_audit
from src.strategy_audit_service import (
    audit_excel, exchange_schedule, fetch_audit_market_bundle, matched_benchmark,
    open_scenarios, prospective_results,
)
from src.strategy_audit_diagnostics import exit_counterfactuals, period_slices, sector_results


def bars():
    days = pd.bdate_range("2026-10-05", periods=5)
    frame = pd.DataFrame(dict(open=[100, 100, 89, 105, 110], high=[101, 101, 92, 110, 112],
                              low=[99, 89, 88, 104, 109], close=[100, 90, 91, 110, 111],
                              fx_open=1., fx_close=1., sma_long=80., reduce_setup=False, sell_setup=False), index=days)
    frame["open_at"] = pd.to_datetime(days, utc=True) + pd.Timedelta(hours=13, minutes=30)
    frame["close_at"] = pd.to_datetime(days, utc=True) + pd.Timedelta(hours=20)
    frame.attrs["session_schedule"] = [dict(session=d.date().isoformat(), open_at=r.open_at.isoformat(), close_at=r.close_at.isoformat()) for d, r in frame.iterrows()]
    frame.attrs["expected_sessions"] = [d.date().isoformat() for d in days]
    frame.attrs["required_last_close_at"] = frame.iloc[-1].close_at.isoformat()
    return frame


def signals():
    return pd.DataFrame([dict(event_id="x", ticker="TEST", signal_session="2026-10-02",
                              observed_at=pd.Timestamp("2026-10-02T21:00Z"), entry_label="Entrada fuerte",
                              entry_score=85, sector="Test")])


def test_authoritative_calendar_accounts_for_us_early_close_and_refuses_unknown():
    schedule = exchange_schedule({"exchangeName": "NMS"}, "2026-11-26", "2026-11-27")
    assert len(schedule) == 1
    assert schedule.iloc[0].market_close.hour == 18
    with pytest.raises(ValueError, match="sin verificar"):
        exchange_schedule({"exchangeName": "UNKNOWN"}, "2026-10-05", "2026-10-09")


def test_calendar_coverage_includes_requested_holidays_and_weekend_not_just_last_session():
    schedule = exchange_schedule({"exchangeName": "NMS"}, "2026-11-26", "2026-11-29")
    assert schedule.index[-1].date().isoformat() == "2026-11-27"
    assert schedule.attrs["calendar_coverage_start"] == "2026-11-26"
    assert schedule.attrs["calendar_coverage_end"] == "2026-11-29"
    with pytest.raises(ValueError, match="no cubre"):
        exchange_schedule({"exchangeName": "NMS"}, "2099-01-01", "2099-01-10")


def test_benchmark_matches_multiple_contributions_not_lump_sum_and_earlier_cutoff():
    frame = bars()
    lots = pd.DataFrame([dict(entry_at=frame.iloc[0].open_at, contribution_usd=100),
                         dict(entry_at=frame.iloc[3].open_at, contribution_usd=200)])
    result = matched_benchmark(lots, frame, "2026-10-09")
    assert result["value_usd"] == pytest.approx((1 + 200 / 105) * 111)
    earlier = matched_benchmark(lots.iloc[:1], frame, "2026-10-07")
    assert earlier["value_usd"] == pytest.approx(91)


def test_benchmark_does_not_value_using_a_stale_close_or_skip_pending_contribution():
    frame = bars()
    lots = pd.DataFrame([dict(entry_at=frame.iloc[-1].open_at, contribution_usd=100)])
    assert matched_benchmark(lots, frame.iloc[:-1], "2026-10-09")["value_usd"] is None
    after = lots.assign(entry_at=pd.Timestamp("2026-10-12T14:00Z"))
    assert matched_benchmark(after, frame, "2026-10-09")["pending_contributions"] == 1
    missing_first = matched_benchmark(lots.assign(entry_at=frame.iloc[0].open_at), frame.iloc[1:], "2026-10-09")
    assert missing_first["pending_contributions"] == 1
    assert missing_first["value_usd"] is None


def test_benchmark_does_not_invent_market_closures_after_calendar_horizon():
    frame = bars()
    lots = pd.DataFrame([dict(entry_at=frame.iloc[0].open_at, contribution_usd=100)])
    assert matched_benchmark(lots, frame, "2026-10-11")["value_usd"] is None
    frame.attrs.update(calendar_coverage_start="2026-10-05", calendar_coverage_end="2026-10-11")
    weekend = matched_benchmark(lots, frame, "2026-10-11")
    assert weekend["value_usd"] == pytest.approx(111)
    assert weekend["data_complete"] is True
    outside = matched_benchmark(lots, frame, "2026-10-12")
    assert outside["value_usd"] is None and outside["data_complete"] is False


def test_benchmark_and_forward_horizon_wait_for_actual_final_bar_availability():
    frame = bars()
    frame["close_available_at"] = frame["close_at"].copy()
    frame.loc[frame.index[-1], "close_available_at"] = pd.Timestamp("2026-10-12T10:00:00Z")
    lots = pd.DataFrame([dict(entry_at=frame.iloc[0].open_at, contribution_usd=100)])
    assert matched_benchmark(lots, frame, "2026-10-09")["value_usd"] is None
    forward = prospective_results(signals(), {"TEST": frame}, end="2026-10-09", horizons=(4, 5))
    assert forward.iloc[0]["status"] == "Observado"
    assert forward.iloc[1]["status"] == "Pendiente"
    assert pd.isna(forward.iloc[1]["return_pct"])


def test_forward_followup_never_counts_missing_session_as_later_session():
    frame = bars()
    good = prospective_results(signals(), {"TEST": frame}, end="2026-10-09", horizons=(5, 20))
    assert good.iloc[0].return_pct == pytest.approx(11)
    assert pd.isna(good.iloc[1].return_pct)
    missing_first = prospective_results(signals(), {"TEST": frame.iloc[1:]}, end="2026-10-09", horizons=(1,))
    missing_middle = prospective_results(signals(), {"TEST": frame.drop(frame.index[1])}, end="2026-10-09", horizons=(3,))
    assert missing_first.iloc[0].status == "Pendiente"
    assert missing_middle.iloc[0].status == "Pendiente"


def test_explicit_fetch_recomputes_units_and_does_not_look_ahead_fx():
    calls = []
    def download(ticker, start, end):
        calls.append((ticker, start, end))
        day = pd.DatetimeIndex(["2026-10-05", "2026-10-06"])
        if ticker.endswith("USD=X"):
            return pd.DataFrame(dict(close=[1.2, 9.]), index=day), {}
        return pd.DataFrame(dict(open=[1000., 1000.], high=[1010., 1010.], low=[990., 990.], close=[1000., 1000.], volume=100), index=day), {"currency": "GBp", "exchangeName": "LSE"}
    bundle = fetch_audit_market_bundle(["TEST.L"], "2026-10-05", "2026-10-06", download=download)
    frame = bundle.bars["TEST.L"]
    assert frame.iloc[-1].open == 10
    assert frame.iloc[-1].fx_open == 1.2
    assert pd.isna(frame.iloc[0].fx_open)
    assert frame.attrs["session_schedule"]
    assert frame.attrs["calendar_coverage_start"] == calls[0][1]
    assert frame.attrs["calendar_coverage_end"] == "2026-10-06"
    assert bundle.manifest["calendars"]["TEST.L"]["calendar_coverage_end"] == "2026-10-06"
    assert calls[0][2] == "2026-10-07"
    assert not bundle.manifest["cash_dividends_included"]


def test_private_excel_has_literal_strings_timezone_safe_dates_and_full_manifest():
    report = run_strategy_audit(signals(), {"TEST": bars()}, AuditConfig(end="2026-10-09", exit_policy="hold"))
    observed = signals().assign(test_string="=HYPERLINK(\"https://example.org\")")
    long_text = "á" * 70000
    content = audit_excel(report, signals=observed, rejected=pd.DataFrame(), input_manifest={"long": long_text})
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["Señales observadas"]
    target = next(cell for row in sheet for cell in row if cell.value == observed.test_string.iloc[0])
    assert target.data_type == "s"
    assert "Lotes" in workbook.sheetnames
    parts = [row for row in workbook["Reglas y procedencia"].iter_rows(min_row=2, values_only=True) if row[0] == "inputs"]
    import json
    assert json.loads("".join(part[2] for part in parts))["long"] == long_text


def test_counterfactual_keeps_loss_and_reports_subsequent_recovery_without_forecast():
    config = AuditConfig(end="2026-10-09", exit_policy="reduce", amount_usd=100)
    report = run_strategy_audit(signals(), {"TEST": bars()}, config)
    hold = run_strategy_audit(signals(), {"TEST": bars()}, replace(config, exit_policy="hold"))
    comparison = exit_counterfactuals(report, hold, {"TEST": bars()}, end="2026-10-09")
    assert report.metrics["realized_pnl_usd"] == pytest.approx(-11)
    assert comparison.iloc[0].missed_terminal_gain_usd == pytest.approx(22)
    assert comparison.iloc[0].avoided_terminal_loss_usd == 0
    assert comparison.iloc[0].sessions_to_breakeven == 2
    assert "not a prediction" in comparison.iloc[0].interpretation
    assert sector_results(hold.lots).iloc[0].roi_pct == pytest.approx(11)
    hypotheses = open_scenarios(hold.lots)
    assert hypotheses.iloc[0].wealth_usd == pytest.approx(111 * 1.05)
    assert "no pronóstico" in hypotheses.iloc[0].scenario


def test_correction_windows_preserve_contributions_and_do_not_delete_bad_days():
    daily = pd.DataFrame(dict(date=["2026-10-05", "2026-10-06", "2026-10-07"], value_usd=[100, 190, 205], contribution_usd=[100, 100, 0]))
    windows = period_slices(daily, ("2026-10-06", "2026-10-06"))
    assert windows.profit_usd.tolist() == [0, -10, 15]
    assert daily.value_usd.tolist() == [100, 190, 205]
    indexed = period_slices(daily.set_index("date"), ("2026-10-06", "2026-10-06"))
    assert indexed.profit_usd.tolist() == [0, -10, 15]


def test_hypotheses_do_not_hide_unresolved_positions_or_unknown_sale_proceeds():
    lots = pd.DataFrame([dict(status="open", value_usd=100.), dict(status="unresolved", value_usd=None)])
    assert open_scenarios(lots).wealth_usd.isna().all()
    assert open_scenarios(lots.iloc[:1], None).wealth_usd.isna().all()
    assert open_scenarios(pd.DataFrame(), 50).wealth_usd.tolist() == [50] * 4
