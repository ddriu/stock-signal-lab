"""Inspect generated export bytes only; no authored workbook or account data."""
from io import BytesIO
import json

from openpyxl import load_workbook

from src.model_results_tracking_ui import export_tracking_excel


def test_tracking_excel_preserves_null_literal_cells_and_complete_event_evidence():
    event = dict(season_id="s", event_id="e", event_type="order_filled",
        occurred_at="2026-10-20T14:00:00Z", available_at="2026-10-20T21:00:00Z",
        owner="private@example.invalid", payload={"ticker": "=SUM(1,2)", "side": "buy",
            "gross_usd": 100, "fx": None, "source": "https://example.invalid",
            "extra": "x" * 40000})
    raw = export_tracking_excel([dict(season_id="s", owner="private@example.invalid", manifest={"mode": "budget"})],
        [dict(season_id="s", snapshot={"coverage_complete": False, "metrics": None,
            "daily": [{"date": "2026-10-20", "value_usd": None}]})], [event])
    book = load_workbook(BytesIO(raw), data_only=False)
    assert {"Estadisticas", "Transacciones", "Evolucion", "Eventos"}.issubset(book.sheetnames)
    assert all(cell.data_type != "f" and cell.hyperlink is None for sheet in book
        for row in sheet for cell in row)
    values = [str(cell.value) for sheet in book for row in sheet for cell in row]
    assert "private@example.invalid" not in " ".join(values)
    archive = book["Eventos"]
    columns = [cell.value for cell in archive[1]]
    content_column = columns.index("content")
    restored = json.loads("".join(row[content_column] for row in archive.iter_rows(min_row=2, values_only=True)))
    assert len(restored["payload"]["extra"]) == 40000
    assert restored["payload"]["fx"] is None
